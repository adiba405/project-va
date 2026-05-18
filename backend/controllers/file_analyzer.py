import os
from io import BytesIO

from flask_jwt_extended import get_jwt_identity
from PIL import Image
import pytesseract
import docx
from pypdf import PdfReader

import openpyxl
import xlrd

import csv

from utils.helpers import parse_object_id


UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), '..', 'uploads')
MAX_TEXT_LENGTH = 10000


def _read_txt(filepath):
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()


def _read_pdf(filepath):
    text = []
    with open(filepath, 'rb') as f:
        reader = PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
    return '\n'.join(text)


def _ocr_pdf_images(filepath):
    text_chunks = []
    with open(filepath, 'rb') as f:
        reader = PdfReader(f)
        for page in reader.pages:
            images = getattr(page, 'images', None) or []
            for image_file in images:
                try:
                    image = Image.open(BytesIO(image_file.data))
                    ocr_text = pytesseract.image_to_string(image)
                    if ocr_text and ocr_text.strip():
                        text_chunks.append(ocr_text)
                except Exception:
                    continue
    return '\n'.join(text_chunks)


def _read_docx(filepath):
    document = docx.Document(filepath)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    return '\n'.join(paragraphs)


def _read_image(filepath):
    image = Image.open(filepath)
    text = pytesseract.image_to_string(image)
    return text or ''


def _read_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        try:
            parts.append(f"Sheet: {ws.title}")
        except Exception:
            parts.append("Sheet")

        row_count = 0
        for row in ws.iter_rows(values_only=True):
            row_count += 1
            if row_count > 2000:
                parts.append("[Truncated rows]")
                break

            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not cells:
                continue

            parts.append(", ".join(cells))

    return "\n".join(parts)


def _read_xls(filepath):
    wb = xlrd.open_workbook(filepath)
    parts = []
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        parts.append(f"Sheet: {sheet_name}")

        max_rows = min(ws.nrows, 2000)
        max_cols = ws.ncols
        for r in range(max_rows):
            row_cells = []
            for c in range(max_cols):
                val = ws.cell_value(r, c)
                if val is None:
                    continue
                s = str(val).strip()
                if s:
                    row_cells.append(s)
            if row_cells:
                parts.append(", ".join(row_cells))

    return "\n".join(parts)


def _normalize_text(text):

    if text is None:
        return ''
    normalized = ' '.join(text.split())
    if len(normalized) > MAX_TEXT_LENGTH:
        return normalized[:MAX_TEXT_LENGTH] + '\n\n[Truncated additional content]'
    return normalized


def _read_csv(filepath):
    # Basic CSV reader (comma/semicolon/tab separated depending on file)
    # Keeps it lightweight; best-effort extraction.
    with open(filepath, 'r', encoding='utf-8', errors='ignore', newline='') as f:
        sample = f.read(8192)
        f.seek(0)

        # crude delimiter detection
        delimiters = [',', ';', '\t', '|']
        delimiter = ','
        best = -1
        for d in delimiters:
            score = sample.count(d)
            if score > best:
                best = score
                delimiter = d

        reader = csv.reader(f, delimiter=delimiter)
        lines = []
        row_count = 0
        for row in reader:
            row_count += 1
            if row_count > 2000:
                lines.append('[Truncated rows]')
                break
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(', '.join(cells))
        return '\n'.join(lines)


def extract_text_from_file(app, file_id):
    user_id = get_jwt_identity()
    file_obj = parse_object_id(file_id, 'file_id')
    if file_obj is None:
        return {'success': False, 'message': 'Invalid file ID', 'status': 400}

    file_doc = app.mongo.db.files.find_one({'_id': file_obj, 'user_id': user_id})
    if not file_doc:
        return {'success': False, 'message': 'File not found', 'status': 404}

    filepath = os.path.join(UPLOAD_FOLDER, file_doc['stored_filename'])
    if not os.path.exists(filepath):
        return {'success': False, 'message': 'File not found on disk', 'status': 404}

    file_type = (file_doc.get('file_type') or '').lower()

    try:
        if file_type == 'txt':
            text = _read_txt(filepath)
        elif file_type == 'pdf':
            text = _read_pdf(filepath)
            if not _normalize_text(text):
                text = _ocr_pdf_images(filepath)
        elif file_type in ('docx', 'doc'):
            text = _read_docx(filepath)
        elif file_type in ('jpg', 'jpeg', 'png'):
            text = _read_image(filepath)
        elif file_type == 'xlsx':
            text = _read_xlsx(filepath)
        elif file_type == 'xls':
            text = _read_xls(filepath)
        else:
            # Last-resort: try treating as plain text
            # (helps if users upload text-like formats without a supported extension)
            try:
                text = _read_txt(filepath)
            except Exception:
                return {'success': False, 'message': 'Unsupported file type for extraction', 'status': 400}
    except Exception as e:
        return {'success': False, 'message': f'Error extracting file text: {str(e)}', 'status': 500}

    normalized = _normalize_text(text)

    if not normalized.strip():
        if file_type == 'pdf':
            return {
                'success': False,
                'message': 'No readable text found. The PDF may be scanned, image-based, or OCR is unavailable.',
                'status': 400
            }
        return {
            'success': False,
            'message': 'No readable text found in the file.',
            'status': 400
        }

    return {'success': True, 'data': {'text': normalized, 'filename': file_doc.get('original_filename')}}
